import os
import io
import re

def extract_text_from_pdf_stream(file_stream_or_bytes):
    """Extract text from PDF in memory (bytes or BytesIO) without touching disk."""
    if isinstance(file_stream_or_bytes, bytes):
        stream = io.BytesIO(file_stream_or_bytes)
    else:
        stream = file_stream_or_bytes
        
    text_content = []
    
    # 1. Try pdfplumber with memory stream
    try:
        import pdfplumber
        stream.seek(0)
        with pdfplumber.open(stream) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_content.append(page_text)
                    
        extracted = "\n\n".join(text_content).strip()
        if extracted:
            return extracted
    except Exception as e:
        print(f"pdfplumber stream error: {e}")
        
    # 2. Fallback to PyPDF2 with memory stream
    try:
        import PyPDF2
        stream.seek(0)
        reader = PyPDF2.PdfReader(stream)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
                
        extracted = "\n\n".join(text_content).strip()
        if extracted:
            return extracted
    except Exception as e:
        print(f"PyPDF2 stream error: {e}")
        
    return "\n\n".join(text_content).strip()

def extract_text_from_docx_stream(file_stream_or_bytes):
    """Extract text from DOCX in memory without touching disk."""
    try:
        import docx
        if isinstance(file_stream_or_bytes, bytes):
            stream = io.BytesIO(file_stream_or_bytes)
        else:
            stream = file_stream_or_bytes
            stream.seek(0)
            
        doc = docx.Document(stream)
        full_text = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
                
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
                    
        return "\n".join(full_text).strip()
    except Exception as e:
        print(f"DOCX stream extraction error: {e}")
        return ""

def extract_text_from_txt_bytes(file_bytes):
    """Extract text from plain text bytes."""
    for encoding in ['utf-8', 'latin-1', 'cp1252']:
        try:
            return file_bytes.decode(encoding).strip()
        except (UnicodeDecodeError, AttributeError):
            continue
    return str(file_bytes)

def clean_extracted_text(raw_text):
    """Normalize whitespace and remove unprintable characters."""
    if not raw_text:
        return ""
    cleaned = raw_text.replace('\x00', ' ')
    cleaned = re.sub(r'\n{3,}', '\n\n', cleaned)
    cleaned = re.sub(r'[ \t]{2,}', ' ', cleaned)
    return cleaned.strip()

def parse_resume_stream(file_storage_or_bytes, filename=""):
    """
    Parses resume text purely in-memory.
    Accepts:
    - Flask FileStorage object (file)
    - io.BytesIO stream
    - raw bytes
    - or file path string
    """
    # 1. Determine filename / extension
    if hasattr(file_storage_or_bytes, 'filename') and file_storage_or_bytes.filename:
        ext = os.path.splitext(file_storage_or_bytes.filename)[1].lower()
    elif filename:
        ext = os.path.splitext(filename)[1].lower()
    elif isinstance(file_storage_or_bytes, str):
        ext = os.path.splitext(file_storage_or_bytes)[1].lower()
    else:
        ext = '.pdf' # default

    # 2. Get bytes or stream
    if hasattr(file_storage_or_bytes, 'read'):
        file_storage_or_bytes.seek(0)
        file_bytes = file_storage_or_bytes.read()
    elif isinstance(file_storage_or_bytes, str) and os.path.exists(file_storage_or_bytes):
        with open(file_storage_or_bytes, 'rb') as f:
            file_bytes = f.read()
    elif isinstance(file_storage_or_bytes, bytes):
        file_bytes = file_storage_or_bytes
    else:
        raise ValueError("Invalid file input provided.")

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    # 3. Extract text based on extension
    raw_text = ""
    if ext == '.pdf':
        raw_text = extract_text_from_pdf_stream(file_bytes)
    elif ext in ['.docx', '.doc']:
        raw_text = extract_text_from_docx_stream(file_bytes)
    elif ext in ['.txt', '.rtf', '.md']:
        raw_text = extract_text_from_txt_bytes(file_bytes)
    else:
        raise ValueError(f"Unsupported format '{ext}'. Please upload a PDF, DOCX, or TXT resume.")

    cleaned = clean_extracted_text(raw_text)
    if not cleaned:
        raise ValueError("Could not extract readable text from the file. Please ensure it is not scanned/image-only or corrupted.")

    return cleaned

# Backwards compatibility helper
def parse_resume_file(file_path_or_storage):
    return parse_resume_stream(file_path_or_storage)
